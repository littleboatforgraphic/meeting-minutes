"""Staff directory import (HR Excel/CSV). The internal DB is the default directory; ERP/M365/Google come later."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .db import Person

HEADERS = {
    "name": "name", "full name": "name", "staff name": "name", "nama": "name",
    "email": "email", "e-mail": "email",
    "phone": "phone", "mobile": "phone", "whatsapp": "phone", "hp": "phone", "tel": "phone",
    "telegram": "telegram_username", "telegram username": "telegram_username",
    "erp": "erp_user_id", "erp user id": "erp_user_id", "erp_user_id": "erp_user_id", "staff id": "erp_user_id",
    "aliases": "aliases", "nickname": "aliases", "nicknames": "aliases",
    "role": "role",
}


def normalize_phone(raw: str | None, country_code: str = "60") -> str | None:
    """Digits only, E.164 without '+'. Malaysian local numbers (01x...) get the country code."""
    if not raw:
        return None
    digits = "".join(c for c in str(raw) if c.isdigit())
    if not digits:
        return None
    if digits.startswith("00"):
        digits = digits[2:]
    elif digits.startswith("0"):
        digits = country_code + digits[1:]
    return digits


def read_rows(data: bytes, filename: str) -> list[dict]:
    if Path(filename).suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = [[("" if c is None else str(c)).strip() for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(data.decode("utf-8-sig")))]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    keys = [HEADERS.get(h.lower().strip()) for h in rows[0]]
    if "name" not in keys:
        raise ValueError("directory file needs a 'Name' column")
    out = []
    for r in rows[1:]:
        rec = {k: v for k, v in zip(keys, r, strict=False) if k and v}
        if rec.get("name"):
            out.append(rec)
    return out


def upsert_people(db: Session, rows: list[dict], country_code: str = "60") -> dict[str, int]:
    stats = {"created": 0, "updated": 0}
    for rec in rows:
        email = (rec.get("email") or "").lower() or None
        phone = normalize_phone(rec.get("phone"), country_code)
        person = None
        if email:
            person = db.scalar(select(Person).where(func.lower(Person.email) == email))
        if not person and phone:
            person = db.scalar(select(Person).where(Person.phone == phone))
        if not person:
            person = db.scalar(select(Person).where(func.lower(Person.name) == rec["name"].lower()))
        if person and email and person.email and person.email.lower() != email:
            person = None  # same name/phone but a different email: a different employee, never overwrite their login
        if not person:
            person = Person(name=rec["name"])
            db.add(person)
            stats["created"] += 1
        else:
            stats["updated"] += 1
        person.name = rec["name"]
        person.email = email or person.email
        person.phone = phone or person.phone
        if rec.get("telegram_username"):
            person.telegram_username = rec["telegram_username"].lstrip("@")
        if rec.get("erp_user_id"):
            person.erp_user_id = rec["erp_user_id"]
        if rec.get("aliases"):
            person.aliases = [a.strip() for a in rec["aliases"].replace(";", ",").split(",") if a.strip()]
        if rec.get("role", "").lower() in ("admin", "member"):
            person.role = rec["role"].lower()
        db.flush()
    return stats
